"""手动复核统一心法名册及满重属性与 DIY 工作簿一致，不参与 CI 收集。"""

from __future__ import annotations

import argparse
import math
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(ROOT / "src"))

from lvjiang.apps.yysls.core.attr_model import AttrModelManager  # noqa: E402

FIELDS = {
    "小外": "min_outer", "大外": "max_outer", "精准": "precision",
    "会心": "crit_rate", "直接会心": "direct_crit", "会意": "intent_rate",
    "直接会意": "direct_intent", "外功穿透": "outer_pen", "外功伤害": "outer_bonus",
    "会心伤害": "crit_dmg", "会意伤害": "intent_dmg",
    "最小鸣金": "min_mingjin", "最大鸣金": "max_mingjin",
    "鸣金穿透": "mingjin_pen", "鸣金加成": "mingjin_bonus",
    "最小裂石": "min_lieshi", "最大裂石": "max_lieshi",
    "裂石穿透": "lieshi_pen", "裂石加成": "lieshi_bonus",
    "最小牵丝": "min_qiansi", "最大牵丝": "max_qiansi",
    "牵丝穿透": "qiansi_pen", "牵丝加成": "qiansi_bonus",
    "最小破竹": "min_pozhu", "最大破竹": "max_pozhu",
    "破竹穿透": "pozhu_pen", "破竹加成": "pozhu_bonus",
    "最小无相": "min_wuxiang", "最大无相": "max_wuxiang", "无相穿透": "wuxiang_pen",
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbook", nargs="?", type=Path,
        default=ROOT / "config/local/data/DIY计算器0.10.1.xlsx",
    )
    workbook = load_workbook(parser.parse_args().workbook, read_only=True, data_only=True)
    try:
        sheet = workbook["心法"]
        headers = next(sheet.iter_rows(values_only=True))
        columns = {index: FIELDS[header] for index, header in enumerate(headers) if header in FIELDS}
        assert set(columns.values()) == set(FIELDS.values()), "心法表列发生变化"
        expected = {}
        for row in sheet.iter_rows(min_row=3, values_only=True):
            if row[0]:
                assert row[0] not in expected, f"表格心法重复: {row[0]}"
                expected[row[0]] = {field: float(row[index] or 0) for index, field in columns.items()}
    finally:
        workbook.close()

    manager = AttrModelManager(ROOT / "config/system/yysls/attr_model")
    assert not manager.errors(), manager.errors()
    effects = manager.effects(("inner_way",))
    assert {effect.group for effect in effects} == set(expected), "心法名册与表格不同"
    for name, values in expected.items():
        entries = [effect for effect in effects if effect.group == name]
        assert sorted(effect.tier for effect in entries) == [1, 2, 3, 4, 5, 6], name
        actual = defaultdict(float)
        for effect in entries:
            assert manager.source_file(effect.source_id) == "inner_way.yaml", effect.source_id
            for field, value in effect.stats.items():
                actual[field] += value
        for field in set(values) | set(actual):
            assert math.isclose(actual[field], values.get(field, 0), abs_tol=1e-9), (
                name, field, actual[field], values.get(field, 0))
    print(f"通过：{len(expected)} 门心法，{len(effects)} 个重数条目，全部满重属性与工作簿一致。")


if __name__ == "__main__":
    main()
