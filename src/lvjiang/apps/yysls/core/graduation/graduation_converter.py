"""Convert graduation workbooks into executable, named JSON schemes."""
from __future__ import annotations

import hashlib
import json
import os
import re
import tempfile
from pathlib import Path
from typing import Any

import openpyxl
import yaml
from openpyxl.worksheet.formula import ArrayFormula

from lvjiang.core.config.resolver import get_resolver
from lvjiang.core.paths import get_config_system_dir

from ..combat.combat_attrs import AFFIX_TO_ATTR
from .excel_formula import FormulaError, FormulaModel, parse_formula
from .graduation_program import ProgramCompiler, ProgramRuntime
from .model_registry import (
    MODEL_SCHEMA_VERSION,
    available_models,
    invalidate_model_registry,
    level_dirname,
    model_filename,
    model_rel_path,
)

GRADUATION_DIR = get_config_system_dir() / "yysls" / "graduation"
GAME_CONFIG_DIR = get_config_system_dir() / "yysls" / "game_config"

#: 第 2~14 行是各流派完全一致的部位属性输入。
INPUTS = {
    "min_outer": "期望!B2", "max_outer": "期望!C2",
    "outer_pen": "期望!D2", "outer_bonus": "期望!E2",
    "min_mingjin": "期望!B3", "max_mingjin": "期望!C3",
    "mingjin_pen": "期望!D3", "mingjin_bonus": "期望!E3",
    "min_lieshi": "期望!B4", "max_lieshi": "期望!C4",
    "lieshi_pen": "期望!D4", "lieshi_bonus": "期望!E4",
    "min_qiansi": "期望!B5", "max_qiansi": "期望!C5",
    "qiansi_pen": "期望!D5", "qiansi_bonus": "期望!E5",
    "min_pozhu": "期望!B6", "max_pozhu": "期望!C6",
    "pozhu_pen": "期望!D6", "pozhu_bonus": "期望!E6",
    "min_wuxiang": "期望!B7", "max_wuxiang": "期望!C7",
    "precision": "期望!B8", "crit_rate": "期望!B9",
    "intent_rate": "期望!B10", "direct_crit": "期望!B11",
    "direct_intent": "期望!B12", "crit_dmg": "期望!B13",
    "intent_dmg": "期望!B14",
}

#: 第 15 行起是装备词条输入区，其行数随表格版本变化：115 起游戏把单体/
#: 群体类奇术增伤合并成全奇术增伤，一部分流派的表格因此少一行，定音从
#: B21 上移到 B20。这整段都不能按固定行号读取。
INPUT_ZONE_START = 15

#: 可以直接作为毕业率方案 field 输入的增伤类别。
BONUS_INPUT_KEYS = {
    "all_skill_bonus", "boss_bonus",
    "single_qs_bonus", "group_qs_bonus", "all_qs_bonus",
}

#: 增伤类别 → 标准词条名。复用战斗属性的权威映射，不再另写一份词条名。
ATTR_AFFIXES: dict[str, set[str]] = {
    field_name: {
        affix for affix, target in AFFIX_TO_ATTR.items()
        if target == field_name
    }
    for field_name in BONUS_INPUT_KEYS
}

RATIO_INPUTS = {
    "outer_bonus", "mingjin_bonus", "lieshi_bonus", "qiansi_bonus",
    "pozhu_bonus", "precision", "crit_rate", "intent_rate", "direct_crit",
    "direct_intent", "crit_dmg", "intent_dmg", "all_skill_bonus",
    "boss_bonus", "single_qs_bonus", "group_qs_bonus", "all_qs_bonus",
}
OUTPUTS = {
    "combat_time": "期望!I8", "total_damage": "期望!I10",
    "dps": "期望!I12", "rdps": "期望!I14",
    "graduation_rate": "期望!I16",
}


def _game_config() -> dict[str, Any]:
    """合并导入阶段用到的游戏配置，去掉只对下发有用的版本号。"""
    config: dict[str, Any] = {}
    for filename in ("affixes.yaml", "equipment.yaml", "schools.yaml"):
        document = yaml.safe_load(
            (GAME_CONFIG_DIR / filename).read_text(encoding="utf-8"),
        ) or {}
        document.pop("content_version", None)
        config.update(document)
    return config


def _alias_index(config: dict[str, Any]) -> dict[str, list[str]]:
    """别名 → 标准词条名。同一别名可以指向多个流派的同类词条。"""
    canonical_names: set[str] = set()
    for category in (config.get("affix_caps") or {}).values():
        raw = category.get("_aliases", []) if isinstance(category, dict) else []
        groups = raw.values() if isinstance(raw, dict) else [raw]
        for names in groups:
            if isinstance(names, list):
                canonical_names.update(str(name) for name in names)
    index: dict[str, list[str]] = {}
    for affix_name, aliases in (config.get("affix_aliases") or {}).items():
        if affix_name not in canonical_names:
            raise RuntimeError(f"affix_aliases uses unknown affix: {affix_name!r}")
        if not isinstance(aliases, list):
            raise RuntimeError(f"affix_aliases[{affix_name!r}] must be a list")
        for alias in aliases:
            index.setdefault(str(alias), []).append(str(affix_name))
    return index


def _school_skill_names(config: dict[str, Any], school: str) -> list[str]:
    """当前流派的指定技能增效词条，按配置顺序。"""
    skill_category = (config.get("affix_caps") or {}).get("指定技能增效") or {}
    skill_groups = skill_category.get("_aliases") if isinstance(
        skill_category, dict,
    ) else None
    if not isinstance(skill_groups, dict):
        raise RuntimeError(
            "游戏配置缺少 affix_caps.指定技能增效._aliases 流派分组，"
            "无法解析 Excel 的指定技能增效。请补充配置后重新导入。"
        )
    school_skill_names = skill_groups.get(school)
    if not isinstance(school_skill_names, list) or not school_skill_names:
        raise RuntimeError(
            f"游戏配置没有流派「{school}」的指定技能增效分组，无法解析 Excel。"
            "请补充流派配置后重新导入。"
        )
    return [str(name) for name in school_skill_names]


def _school_weapon_affixes(config: dict[str, Any], school: str) -> set[str]:
    """当前流派主副武器对应的武学增效词条。"""
    school_config = (config.get("schools") or {}).get(school)
    if not isinstance(school_config, dict):
        raise RuntimeError(f"游戏配置缺少流派「{school}」定义，无法筛选 Excel 词条。")
    weapon_affixes = {
        str(entry.get("name")): str(entry.get("wuxue_affix"))
        for entry in (config.get("weapon_types") or [])
        if isinstance(entry, dict) and entry.get("name") and entry.get("wuxue_affix")
    }
    resolved: set[str] = set()
    for position in ("main", "sub"):
        affix = weapon_affixes.get(
            str((school_config.get(position) or {}).get("weapon")))
        if affix is not None:
            resolved.add(affix)
    return resolved


def _affix_input_key(
    name: str, weapon_affixes: set[str], skill_names: set[str],
) -> str | None:
    """一个标准词条名属于哪一类毕业率输入。"""
    if name in skill_names:
        return "special_bonus"
    if name in weapon_affixes:
        return "weapon_bonus"
    field_name = AFFIX_TO_ATTR.get(name)
    return field_name if field_name in BONUS_INPUT_KEYS else None


def _input_key(
    label: str, row: int, alias_index: dict[str, list[str]],
    weapon_affixes: set[str], skill_names: set[str],
) -> str:
    """把一行 Excel 简称解析成输入类别。"""
    candidates = alias_index.get(label) or []
    keys = {
        key for name in candidates
        if (key := _affix_input_key(name, weapon_affixes, skill_names)) is not None
    }
    if len(keys) == 1:
        return keys.pop()
    if not candidates:
        raise RuntimeError(
            f"Excel 期望页第 {row} 行的输入简称 {label!r} 不在词组别名配置中。"
            "请在 affix_aliases 中补充该简称后重新导入。"
        )
    raise RuntimeError(
        f"Excel 期望页第 {row} 行的输入简称 {label!r} 应唯一映射到一类"
        f"毕业率输入，实际匹配 {sorted(keys) or '无'}（候选词条 {candidates!r}）。"
        "请修正 Excel 或词组别名配置后重新导入。"
    )


def _next_special_bonus_key(inputs: dict[str, str]) -> str:
    """首个定音沿用 ``special_bonus``，后续依次编号以维持旧方案结构。"""
    if "special_bonus" not in inputs:
        return "special_bonus"
    index = 2
    while f"special_bonus_{index}" in inputs:
        index += 1
    return f"special_bonus_{index}"


def _workbook_inputs(
    workbook, school: str, config: dict[str, Any],
) -> dict[str, str]:
    """返回该工作簿的输入单元格。

    第 2~14 行是各流派一致的部位属性；第 15 行起到「固伤加成」行之前是装备
    词条，行数随表格版本变化——115 起单体/群体类奇术增伤合并成全奇术增伤，
    一部分流派因此少一行，定音从 B21 上移到 B20，而另一部分流派仍沿用两行
    奇术。这整段都必须按 A 列简称反查词条类别定位，不能按行号猜。
    """
    sheet = workbook["期望"]
    fixed_row = _label_row(sheet, "固伤加成")
    alias_index = _alias_index(config)
    weapon_affixes = _school_weapon_affixes(config, school)
    skill_names = set(_school_skill_names(config, school))
    inputs = dict(INPUTS)
    for row in range(INPUT_ZONE_START, fixed_row):
        label = str(sheet.cell(row, 1).value or "").strip()
        if not label:
            continue
        key = _input_key(
            label, row, alias_index, weapon_affixes, skill_names)
        if key == "weapon_bonus":
            key = ("weapon_bonus_primary" if "weapon_bonus_primary" not in inputs
                   else "weapon_bonus_secondary")
        elif key == "special_bonus":
            key = _next_special_bonus_key(inputs)
        if key in inputs:
            raise RuntimeError(
                f"Excel 期望页第 {row} 行的简称 {label!r} 与 B{inputs[key].rsplit('B', 1)[-1]}"
                f" 同属 {key} 输入，无法确定应以哪一行为准。"
            )
        inputs[key] = f"期望!B{row}"
    return inputs


def _is_dynamic_affix_input(name: str) -> bool:
    return name in {"weapon_bonus_primary", "weapon_bonus_secondary"} or (
        name == "special_bonus" or name.startswith("special_bonus_")
    )


def _json_value(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _version(filename: str) -> str:
    match = re.search(r"(\d+\.\d+)\.xlsx$", filename)
    return match.group(1) if match else "unknown"


def _affix_input_names(
    workbook, school: str, inputs: dict[str, str], config: dict[str, Any],
) -> dict[str, list[str]]:
    """把词条类输入的标准词条名解析出来。

    ``_workbook_inputs`` 已经按简称确定每一行属于哪一类输入，这里只负责在
    该类别的候选词条里做唯一映射；输入名必须进报错信息，否则用户只会看到
    一个简称匹配失败，无法判断是哪一行读错了。
    """
    named_inputs = {
        "all_skill_bonus", "boss_bonus", "weapon_bonus_primary",
        "weapon_bonus_secondary", "single_qs_bonus", "group_qs_bonus",
        "all_qs_bonus",
    }
    labels: dict[str, Any] = {}
    for name, target in inputs.items():
        if name not in named_inputs and not _is_dynamic_affix_input(name):
            continue
        _sheet, coordinate = target.split("!", 1)
        row = workbook["期望"][coordinate].row
        labels[name] = workbook["期望"].cell(row, 1).value
    alias_index = _alias_index(config)
    school_skill_names = _school_skill_names(config, school)
    school_weapon_names = _school_weapon_affixes(config, school)
    allowed_by_input = {
        "all_skill_bonus": ATTR_AFFIXES["all_skill_bonus"],
        "boss_bonus": ATTR_AFFIXES["boss_bonus"],
        "weapon_bonus_primary": school_weapon_names,
        "weapon_bonus_secondary": school_weapon_names,
        "single_qs_bonus": ATTR_AFFIXES["single_qs_bonus"],
        "group_qs_bonus": ATTR_AFFIXES["group_qs_bonus"],
        "all_qs_bonus": ATTR_AFFIXES["all_qs_bonus"],
    }
    result: dict[str, list[str]] = {}
    for input_name, label in labels.items():
        if label is None or not str(label).strip():
            if input_name == "special_bonus":
                # 部分流派没有指定技能增效；Excel 留空即明确表示不参与计算。
                result[input_name] = []
                continue
            raise RuntimeError(
                f"{school} Excel 输入 {input_name} 的字段名为空。"
                "请修正 Excel，或在词组配置中补充对应别名后重新导入。"
            )
        if input_name == "special_bonus" or input_name.startswith(
            "special_bonus_"
        ):
            # 只在当前流派的指定技能增效分组中反查，按配置顺序取第一个。
            matches = [
                str(affix_name) for affix_name in school_skill_names
                if str(label) in (config.get("affix_aliases") or {}).get(
                    str(affix_name), [],
                )
            ]
            if not matches:
                raise RuntimeError(
                    f"流派「{school}」无法解析 Excel 输入 {input_name} 的"
                    f"指定技能增效简称 {label!r}。请在该流派的指定技能增效"
                    "词条中配置此别名，或修正 Excel 字段名后重新导入。"
                )
            result[input_name] = [matches[0]]
            continue
        matches = [
            name for name in alias_index.get(str(label), [])
            if name in allowed_by_input[input_name]
        ]
        if len(matches) != 1:
            raise RuntimeError(
                f"流派「{school}」的 Excel 输入 {input_name}（简称 {label!r}）"
                f"应唯一映射到一个精准词条，实际匹配 {matches!r}。"
                "请修改 Excel 或词组别名配置后重新导入。"
            )
        result[input_name] = matches
    return result


def _label_ref(target: str) -> str:
    """输入的简称就在同一行的 A 列。"""
    sheet, coordinate = target.split("!", 1)
    return f"{sheet}!A{coordinate.lstrip('ABCDEFGHIJKLMNOPQRSTUVWXYZ')}"


def _cell_value(formulas, cached, address: str) -> Any:
    sheet, coordinate = address.split("!", 1)
    cached_value = cached[sheet][coordinate].value
    return _json_value(cached_value if cached_value is not None else formulas[sheet][coordinate].value)


def _label_row(sheet, label: str) -> int:
    row = next(
        (index for index in range(1, sheet.max_row + 1)
         if str(sheet.cell(index, 1).value or "").strip() == label),
        None,
    )
    if row is None:
        raise RuntimeError(f"Excel 期望页缺少「{label}」行")
    return row


def _extract_environment(formulas, cached) -> dict[str, Any]:
    sheet = formulas["期望"]
    cached_sheet = cached["期望"]
    fixed_row = _label_row(sheet, "固伤加成")
    food_row = _label_row(sheet, "食物加成")
    team_buffs = []
    for row in range(16, 21):
        for name_col, enabled_col in ((3, 4), (5, 6)):
            name = sheet.cell(row, name_col).value
            if name not in (None, ""):
                team_buffs.append({
                    "name": str(name),
                    "enabled": str(sheet.cell(row, enabled_col).value or "") == "√",
                })
    monster = {}
    for row in range(2, 6):
        name = sheet.cell(row, 8).value
        if name not in (None, ""):
            value = cached_sheet.cell(row, 9).value
            if value is None:
                value = sheet.cell(row, 9).value
            monster[str(name)] = _json_value(value)
    return {
        "food_bonus": {
            "min_outer": _cell_value(
                formulas, cached, f"期望!B{food_row}"),
            "max_outer": _cell_value(
                formulas, cached, f"期望!B{food_row + 1}"),
        },
        "fixed_damage_bonus": _cell_value(
            formulas, cached, f"期望!B{fixed_row}"),
        "team_buffs": team_buffs,
        "monster": monster,
        "combat_time": _cell_value(formulas, cached, "期望!I8"),
    }


def _baseline_attrs(
    model: dict[str, Any], affix_names: dict[str, list[str]],
    inputs: dict[str, str],
) -> dict[str, Any]:
    runtime = FormulaModel(model)
    attrs: dict[str, Any] = {}
    extra_attrs: dict[str, float] = {}
    for name, target in inputs.items():
        value = float(runtime.value(target) or 0)
        if _is_dynamic_affix_input(name):
            names = affix_names[name]
            if names:
                extra_attrs[names[0]] = value
        else:
            attrs[name] = value
    if extra_attrs:
        attrs["extra_attrs"] = extra_attrs
    return attrs


def _compile_v2(
    workbook_model: dict[str, Any], school: str, affix_names: dict[str, list[str]],
    environment: dict[str, Any], inputs: dict[str, str],
) -> dict[str, Any]:
    bindings: dict[str, dict[str, str]] = {}
    for name, target in inputs.items():
        if _is_dynamic_affix_input(name):
            names = affix_names[name]
            if not names:
                # A deliberately empty skill bonus remains a fixed zero.
                continue
            bindings[target] = {"kind": "affix", "name": names[0]}
        else:
            bindings[target] = {"kind": "field", "name": name}
    compiler = ProgramCompiler(workbook_model, bindings)
    program = compiler.compile({
        name: OUTPUTS[name]
        for name in ("combat_time", "total_damage", "dps")
    })
    baseline_attrs = _baseline_attrs(workbook_model, affix_names, inputs)
    input_values = []
    extra = baseline_attrs.get("extra_attrs", {})
    for spec in program["inputs"]:
        input_values.append(float(
            baseline_attrs.get(spec["name"], 0)
            if spec["kind"] == "field" else extra.get(spec["name"], 0)
        ))
    compiled_outputs = ProgramRuntime(program, input_values).outputs()
    workbook_runtime = FormulaModel(workbook_model)
    reference = {
        name: float(workbook_runtime.value(OUTPUTS[name]))
        for name in ("combat_time", "total_damage", "dps", "graduation_rate")
    }
    for name, actual in reference.items():
        sheet, coordinate = OUTPUTS[name].split("!", 1)
        cached = workbook_model["sheets"][sheet]["cells"][coordinate].get("cached")
        if cached is not None and abs(actual - float(cached)) > max(
            1e-6, abs(float(cached)) * 1e-10,
        ):
            raise FormulaError(
                f"Excel 公式 {name}={actual} 与工作簿缓存值 {cached} 不一致"
            )
    for name, actual in compiled_outputs.items():
        expected = reference[name]
        if abs(actual - expected) > max(1e-6, abs(expected) * 1e-10):
            raise FormulaError(
                f"编译后的 {name}={actual} 与 Excel 公式结果 {expected} 不一致"
            )
    if reference["graduation_rate"] <= 0:
        raise FormulaError("Excel 的毕业率必须大于 0，无法提取 100% 基准 DPS")
    graduation_baseline_dps = (
        reference["dps"] / reference["graduation_rate"]
    )
    return {
        "content_version": 1,
        "schema_version": MODEL_SCHEMA_VERSION,
        "school": school,
        "source": workbook_model["model"]["source"],
        "baseline_attrs": baseline_attrs,
        "environment": environment,
        "reference": reference,
        "graduation_baseline_dps": graduation_baseline_dps,
        "program": program,
    }


def convert_workbook(
    path: Path,
    school: str,
    scheme: str = "基础方案",
    model_level: int = 110,
    model_version: int = 1,
) -> dict[str, Any]:
    formulas = openpyxl.load_workbook(path, data_only=False, read_only=False)
    cached = openpyxl.load_workbook(path, data_only=True, read_only=False)
    try:
        config = _game_config()
        inputs = _workbook_inputs(formulas, school, config)
        affix_names = _affix_input_names(formulas, school, inputs, config)
        sheets: dict[str, Any] = {}
        formula_count = 0
        for worksheet in formulas.worksheets:
            cached_sheet = cached[worksheet.title]
            cells: dict[str, Any] = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    value = cell.value
                    if value is None:
                        continue
                    formula = value.text if isinstance(value, ArrayFormula) else value
                    if isinstance(formula, str) and formula.startswith("="):
                        # 位置必须进异常：原来只报公式原文，用户拿着一串
                        # `unsupported formula syntax` 既不知道是哪张表哪一格，
                        # 也会以为是律匠不支持这种写法而不是表本身坏了。
                        try:
                            parse_formula(formula)
                        except FormulaError as exc:
                            raise FormulaError(
                                f"{worksheet.title}!{cell.coordinate} "
                                f"的公式无法解析：{exc}（公式：{formula}）"
                            ) from exc
                        entry: dict[str, Any] = {"formula": formula}
                        cached_value = cached_sheet[cell.coordinate].value
                        if cached_value is not None:
                            entry["cached"] = _json_value(cached_value)
                        formula_count += 1
                    else:
                        entry = {"value": _json_value(value)}
                    cells[cell.coordinate] = entry
            sheets[worksheet.title] = {
                "dimensions": {
                    "max_row": worksheet.max_row,
                    "max_column": worksheet.max_column,
                },
                "cells": cells,
            }
        source_bytes = path.read_bytes()
        workbook_model = {
            "schema_version": "1.0",
            "formula_language": "excel_subset_v1",
            "model": {
                "id": school.replace("·", ""), "school": school,
                "source": {
                    "file": path.name, "version": _version(path.name),
                    "sha256": hashlib.sha256(source_bytes).hexdigest(),
                },
                "formula_count": formula_count,
            },
            "inputs": {
                name: {
                    "type": "number",
                    "unit": "ratio" if (
                        name in RATIO_INPUTS or _is_dynamic_affix_input(name)
                    ) else "point",
                    "target": target,
                    **({"label_ref": _label_ref(target),
                        "affix_names": affix_names[name]}
                       if name in affix_names else {}),
                }
                for name, target in inputs.items()
            },
            "outputs": {name: {"ref": ref} for name, ref in OUTPUTS.items()},
            "sheets": sheets,
        }
        model = _compile_v2(
            workbook_model, school, affix_names,
            _extract_environment(formulas, cached), inputs,
        )
        model["scheme"] = scheme
        model["model_level"] = model_level
        model["model_version"] = model_version
        return model
    finally:
        formulas.close()
        cached.close()


def validate_model(model: dict[str, Any]) -> dict[str, float]:
    if model.get("schema_version") != MODEL_SCHEMA_VERSION:
        raise FormulaError(
            f"graduation model must use schema version {MODEL_SCHEMA_VERSION}")
    baseline = model["baseline_attrs"]
    extra = baseline.get("extra_attrs", {})
    values = [
        float(baseline.get(spec["name"], 0) if spec["kind"] == "field"
              else extra.get(spec["name"], 0))
        for spec in model["program"]["inputs"]
    ]
    results = ProgramRuntime(model["program"], values).outputs()
    baseline_dps = float(model["graduation_baseline_dps"])
    if baseline_dps <= 0:
        raise FormulaError("100%毕业率基准 DPS 必须大于 0")
    results["graduation_rate"] = results["dps"] / baseline_dps
    for name, expected in model["reference"].items():
        if abs(results[name] - float(expected)) > max(1e-6, abs(float(expected)) * 1e-10):
            raise FormulaError(f"compiled {name} does not match its reference value")
    return results


def scheme_path(
    school: str, scheme: str, level: int, version: int,
) -> Path:
    return (GRADUATION_DIR / level_dirname(level)
            / model_filename(school, scheme, version))


def write_model(path: Path, model: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = json.dumps(model, ensure_ascii=False, indent=2) + "\n"
    fd, temp_name = tempfile.mkstemp(dir=path.parent, suffix=".json.tmp")
    temp_path = Path(temp_name)
    try:
        with os.fdopen(fd, "w", encoding="utf-8") as stream:
            stream.write(payload)
        os.replace(temp_path, path)
    except BaseException:
        temp_path.unlink(missing_ok=True)
        raise


def import_graduation_scheme(
    excel_path: str | Path, school: str, scheme: str, model_level: int,
) -> tuple[Path, dict[str, float]]:
    source = Path(excel_path)
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("请选择有效的 .xlsx 文件")
    versions = [
        ref.version for ref in available_models(school, scheme)
        if ref.level == model_level
    ]
    model_version = max(versions, default=0) + 1
    model = convert_workbook(
        source, school, scheme, model_level, model_version)
    outputs = validate_model(model)
    rel_path = model_rel_path(school, scheme, model_level, model_version)
    payload = json.dumps(model, ensure_ascii=False, indent=2) + "\n"
    destination = get_resolver().write_entity(rel_path, payload)
    invalidate_model_registry()
    return destination, outputs
